#!/usr/bin/env python3
"""Publica un overlay incremental de texto acreditado para reactivos.

Lee únicamente metadatos ya indexados y documentación pública. Nunca abre
valores de microdatos. Las tablas históricas conservan orden y citas por fila;
cada fila del overlay apunta a su ``id_origen`` estable (fuente + posición).
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import sys
import tempfile
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
RAW_ROOT = REPO_ROOT / "data" / "raw"
OUT_PATH = REPO_ROOT / "data" / "inventario-reactivos-contexto-v1_1.tsv"
PREVIOUS_OVERLAY_PATH = REPO_ROOT / "data" / "inventario-reactivos-contexto-v1_0.tsv"
RESIDUAL_PATH = REPO_ROOT / "data" / "reactivos-contexto-residual-v1_1.tsv"
SOURCES_PATH = REPO_ROOT / "data" / "reactivos-contexto-fuentes-v1_0.tsv"
VERIFIED_PATH = REPO_ROOT / "data" / "reactivos-contexto-verificados-v1_0.tsv"
FD_EXT_PATH = REPO_ROOT / "data" / "inventario-fd-ext-v1_0.tsv"
CACHE_DIR = REPO_ROOT / "data" / ".reactivos-contexto-cache"
EXTRACTOR_VERSION = "reactivos-contexto-1.1.4"

METADATA_SOURCES = {
    "v1_2": REPO_ROOT / "data" / "inventario-reactivos-v1_2.tsv",
    "ext": REPO_ROOT / "data" / "inventario-reactivos-ext-v1_0.tsv",
}
PRIORITY = ("envipe", "ennvih", "encuci", "enif", "ensafi")

OUT_FIELDS = [
    "id_origen", "payload_id", "sha256_12", "instrumento", "ola",
    "archivo_miembro", "variable_id", "texto_reactivo", "texto_tipo",
    "contexto_busqueda", "metodo", "universo_declarado", "fuente_texto",
    "fuente_sha256_12", "referencia_fuente",
]
RESIDUAL_FIELDS = [
    "instrumento", "ola", "payload_id", "archivo_miembro", "motivo",
    "filas_fisicas", "variables_logicas", "variables_ejemplo", "siguiente_accion",
]

ID_LABELS = {"mnemonico", "nemonico", "nombre de la columna", "nombre", "clave", "variable"}
TEXT_LABELS = {"pregunta y categoria", "pregunta", "descripcion del contenido del campo",
               "descripcion", "etiqueta", "observaciones"}
TABLE_IDENTITIES = (
    ("encuci2020sec910", "ENCUCI_2020_SEC_9_10"),
    ("encuci2020sec678", "ENCUCI_2020_SEC_6_7_8"),
    ("encuci2020sec45", "ENCUCI_2020_SEC_4_5"),
    ("encuci2020sd", "ENCUCI_2020_SD"),
    ("encuci2020cs", "ENCUCI_2020_SD"),
    ("encuci2020viv", "ENCUCI_2020_VIV"),
    ("tpervic1", "TPer_Vic1"),
    ("tpervic2", "TPer_Vic2"),
    ("tpervic", "TPer_Vic"),
    ("tperviv", "TPer_Vic"),  # errata del FD ENVIPE 2012
    ("tmodvic", "TMod_Vic"),
    ("tviviendas", "TVivienda"),
    ("tvivienda", "TVivienda"),
    ("thogar", "THogar"),
    ("tsdem", "TSDem"),
    ("tmodulo1", "TModulo1"),
    ("tmodulo2", "TModulo2"),
    ("tmodulo3", "TModulo3"),
    ("tmodulo", "TModulo"),
)


def sanitize(value) -> str:
    return " ".join(str(value or "").replace("\t", " ").split())


def fold(value) -> str:
    text = unicodedata.normalize("NFKD", sanitize(value).lower())
    return "".join(char for char in text if not unicodedata.combining(char))


def compact(value) -> str:
    return re.sub(r"[^a-z0-9]", "", fold(value))


def variable_key(value: str) -> tuple[str | int, ...]:
    return tuple(int(part) if part.isdigit() else part.lower()
                 for part in re.split(r"_+", value) if part)


def resolve_variable(value: str, variables: set[str]) -> str | None:
    lowered = value.lower()
    if lowered in variables:
        return lowered
    key = variable_key(lowered)
    matches = [candidate for candidate in variables if variable_key(candidate) == key]
    return matches[0] if len(matches) == 1 else None


def clean_pdf_text(value: str) -> str:
    text = sanitize(value)
    text = re.split(
        r"\s+(?:\(?contin[uú]a\)?|cons\.\s+pregunta|inegi\.\s+encuesta\s+nacional)",
        text, maxsplit=1, flags=re.I)[0]
    return sanitize(text)


def valid_data_type(value: str) -> bool:
    normalized = fold(value)
    return any(stem in normalized for stem in
               ("alfanumeric", "numeric", "caracter", "texto", "cadena", "string"))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_tsv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as handle:
        return list(csv.DictReader((line for line in handle if not line.startswith("#")), delimiter="\t"))


def instrument_for(row: dict) -> str:
    current = row["instrumento"].lower()
    payload = row["payload_id"].lower()
    if payload.startswith("ennvih/"):
        match = re.search(r"(?:ehh|eloc)(02|05|09)", payload)
        return f"ennvih20{match.group(1)}" if match else "ennvih"
    return current


def year_for(instrument: str, current: str) -> str:
    match = re.search(r"(19|20)\d{2}", instrument)
    return match.group(0) if match else current


def selected(instrument: str, payload: str, objects: list[str]) -> bool:
    haystack = fold(f"{instrument} {payload}")
    return any(fold(obj) in haystack for obj in objects)


def find_header(values: tuple) -> tuple[int | None, int | None, int | None]:
    normalized = [fold(v).rstrip(":.") for v in values]
    id_col = next((i for i, value in enumerate(normalized) if value in ID_LABELS), None)
    text_col = next((i for i, value in enumerate(normalized) if value in TEXT_LABELS), None)
    concept_col = next((i for i, value in enumerate(normalized)
                        if value in {"concepto", "etiqueta de variable"}), None)
    return id_col, text_col, concept_col


def extract_sheet(sheet, source: str, member: str) -> list[dict]:
    result = []
    id_col = text_col = concept_col = None
    current_question = ""
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
        new_id, new_text, new_concept = find_header(values)
        if new_id is not None and new_text is not None:
            id_col, text_col, concept_col, current_question = (
                new_id, new_text, new_concept, "")
            continue
        if id_col is None or text_col is None:
            continue
        variable = sanitize(values[id_col] if id_col < len(values) else "")
        text = sanitize(values[text_col] if text_col < len(values) else "")
        if text and not variable:
            if not re.fullmatch(r"[\(\[]\d+[\)\]]", text):
                current_question = text
            continue
        if not variable or re.fullmatch(r"[\(\[]\d+[\)\]]", variable):
            continue
        if text:
            current_question = text
        item = ""
        if text_col + 1 < id_col:
            item = sanitize(" ".join(sanitize(value) for value in values[text_col + 1:id_col]))
        concept = sanitize(values[concept_col]) if concept_col is not None and concept_col < len(values) else ""
        accredited = sanitize(" ".join(part for part in (current_question, item) if part))
        text_type = "PREGUNTA_DICCIONARIO"
        if not accredited and concept:
            accredited = concept
            text_type = "DESCRIPCION_ADMINISTRATIVA"
        if accredited:
            result.append({
                "variable_id": variable,
                "texto_reactivo": accredited,
                "texto_tipo": text_type,
                "contexto_busqueda": "",
                "fuente_texto": source,
                "referencia_fuente": f"miembro={member};hoja={sheet.title};fila={row_number}",
                "tabla_documental": sheet.title,
            })
    return result


def extract_xlsx(content_or_path, source: str, member: str | None = None) -> list[dict]:
    from openpyxl import load_workbook
    workbook = load_workbook(content_or_path, read_only=True, data_only=True)
    try:
        return [item for sheet in workbook.worksheets
                for item in extract_sheet(sheet, source, member or source)]
    finally:
        workbook.close()


def extract_xls(path: Path, source: str) -> list[dict]:
    import xlrd
    workbook = xlrd.open_workbook(str(path))
    result = []
    for raw_sheet in workbook.sheets():
        class SheetAdapter:
            title = raw_sheet.name

            @staticmethod
            def iter_rows(values_only=True):
                return (tuple(raw_sheet.row_values(index)) for index in range(raw_sheet.nrows))

        result.extend(extract_sheet(SheetAdapter(), source, source))
    return result


def page_map(path: Path, member: str | None, variables: set[str]) -> dict[str, int]:
    import pdfplumber
    if member:
        with zipfile.ZipFile(path) as archive:
            handle = io.BytesIO(archive.read(member))
        pdf_context = pdfplumber.open(handle)
    else:
        pdf_context = pdfplumber.open(str(path))
    found = {}
    with pdf_context as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            for variable in variables - found.keys():
                if re.search(rf"(?<![A-Za-z0-9_]){re.escape(variable)}(?![A-Za-z0-9_])", text, re.I):
                    found[variable] = page_number
    return found


def extract_fd_pdf_index(path: Path, source: str, instrument: str) -> list[dict]:
    indexed = [row for row in read_tsv(FD_EXT_PATH)
               if row["payload_id"] == source and row["instrumento"].lower() == instrument.lower()]
    by_member: dict[str, list[dict]] = defaultdict(list)
    for row in indexed:
        by_member[row["archivo_miembro"]].append(row)
    result = []
    for member, rows in by_member.items():
        zip_member = member if path.suffix.lower() == ".zip" else None
        located = page_map(path, zip_member, {row["variable_id"] for row in rows})
        for row in rows:
            page = located.get(row["variable_id"])
            if not page:
                continue
            result.append({
                "variable_id": row["variable_id"],
                "texto_reactivo": row["texto_reactivo"],
                "texto_tipo": "PREGUNTA_DICCIONARIO",
                "contexto_busqueda": "",
                "fuente_texto": source,
                "referencia_fuente": f"miembro={member};pagina={page}",
                "tabla_documental": member,
            })
    return result


def metadata_tables(instrument: str) -> tuple[set[str], dict[str, set[str]]]:
    """Devuelve variables y tablas históricas del instrumento, sin abrir valores."""
    variables = set()
    tables: dict[str, set[str]] = defaultdict(set)
    for metadata_path in METADATA_SOURCES.values():
        for row in read_tsv(metadata_path):
            if instrument_for(row).lower() != instrument.lower():
                continue
            variable = row["variable_id"].lower()
            variables.add(variable)
            tables[variable].add(row["archivo_miembro"])
    return variables, tables


def table_identity(value: str) -> str | None:
    normalized = compact(value)
    for token, identity in TABLE_IDENTITIES:
        if token in normalized:
            return identity
    return None


def table_matches(member: str, documented: str) -> bool:
    left, right = compact(member), compact(documented)
    left_identity, right_identity = table_identity(left), table_identity(right)
    if left_identity or right_identity:
        return left_identity == right_identity
    return min(len(left), len(right)) >= 5 and (left in right or right in left)


def inferred_table(possible_tables: set[str]) -> str:
    identities = {identity for table in possible_tables
                  if (identity := table_identity(table)) is not None}
    return next(iter(identities)) if len(identities) == 1 else ""


def _line_groups(words: list[dict]) -> list[list[dict]]:
    groups: list[list[dict]] = []
    for word in sorted((word for word in words if word.get("upright", True)),
                       key=lambda item: (item["top"], item["x0"])):
        if not groups or abs(groups[-1][0]["top"] - word["top"]) > 2.0:
            groups.append([word])
        else:
            groups[-1].append(word)
    return groups


def _line_text(words: list[dict], x_max: float | None = None) -> str:
    selected = [word for word in words if x_max is None or word["x0"] < x_max]
    value = sanitize(" ".join(word["text"] for word in sorted(selected, key=lambda item: item["x0"])))
    return re.sub(r"^\d+\s+(?=\d+[.)]|[A-Za-zÁÉÍÓÚÜÑ¿])", "", value)


def _question_stem(variable_id: str) -> str:
    parts = variable_id.split("_")
    return "_".join(parts[:-1]) if len(parts) >= 3 else variable_id


def extract_pdf_words(page, variables: set[str], variable_tables: dict[str, set[str]],
                      source: str, document_name: str, page_number: int,
                      current_table: str) -> list[dict]:
    """Respaldo geométrico para páginas cuyas tablas no tienen bordes."""
    words = page.extract_words(x_tolerance=1, y_tolerance=2, keep_blank_chars=False)
    hits = []
    for word in words:
        token = resolve_variable(word["text"].strip(".,;:"), variables)
        if token is None:
            continue
        if not re.match(r"^(?:[ab]?p)\d", token, re.I):
            continue
        right = [candidate for candidate in words
                 if abs(candidate["top"] - word["top"]) <= 2.0
                 and word["x1"] < candidate["x0"] < word["x1"] + 155]
        if any(valid_data_type(candidate["text"]) for candidate in right):
            hits.append({**word, "variable_id": token})

    rows: list[list[dict]] = []
    for hit in sorted(hits, key=lambda item: (item["top"], item["x0"])):
        if not rows or abs(rows[-1][0]["top"] - hit["top"]) > 2.0:
            rows.append([hit])
        else:
            rows[-1].append(hit)
    lines = _line_groups(words)
    result = []
    group_question = ""
    group_stem = ""
    for index, row in enumerate(rows):
        top = min(item["top"] for item in row)
        bottom = max(item["bottom"] for item in row)
        variable_x = min(item["x0"] for item in row)
        previous_bottom = rows[index - 1][0]["bottom"] if index else max(0, top - 90)
        next_top = rows[index + 1][0]["top"] if index + 1 < len(rows) else min(page.height, bottom + 65)

        before = [_line_text(line) for line in lines
                  if previous_bottom + 0.1 < line[0]["top"] < top - 2
                  and any(word["x0"] < variable_x for word in line)]
        starts = [position for position, value in enumerate(before)
                  if re.match(r"^\d+(?:\.\d+)+[a-z.]?\s", value, re.I)]
        heading = clean_pdf_text(" ".join(before[starts[-1]:])) if starts else ""

        item_candidates = [(line[0]["top"], _line_text(line, variable_x - 2))
                           for line in lines
                           if top - 2 <= line[0]["top"] < next_top - 2
                           and any(word["x0"] < variable_x - 2 for word in line)]
        item_lines = ([value for _, value in item_candidates]
                      if item_candidates and item_candidates[0][0] <= bottom + 2 else [])
        next_heading = next((position for position, value in enumerate(item_lines)
                             if re.match(r"^\d+(?:\.\d+)+[a-z.]?\s", value, re.I)), None)
        if next_heading is not None:
            item_lines = item_lines[:next_heading]
        item = clean_pdf_text(" ".join(value for value in item_lines
                                       if value and not re.fullmatch(r"\d+", value)
                                       and fold(value) not in {"pregunta", "(1)", "continua"}))
        stem = _question_stem(row[0]["variable_id"])
        if heading:
            group_question, group_stem = heading, stem
        elif stem != group_stem:
            group_question, group_stem = "", stem
        question = clean_pdf_text(" ".join(dict.fromkeys(
            part for part in (group_question, item) if part)))
        if not question:
            continue
        for hit in row:
            possible_tables = variable_tables[hit["variable_id"]]
            documented_table = current_table
            if current_table and not any(table_matches(table, current_table)
                                         for table in possible_tables):
                documented_table = ""
            if not documented_table:
                documented_table = inferred_table(possible_tables)
            result.append({
                "variable_id": hit["variable_id"],
                "texto_reactivo": question,
                "texto_tipo": ("PREGUNTA_DICCIONARIO" if group_question or "?" in question
                               else "DESCRIPCION_ADMINISTRATIVA"),
                "contexto_busqueda": "",
                "fuente_texto": source,
                "referencia_fuente": (
                    f"miembro={document_name};pagina={page_number};tabla={documented_table or 'NO_DETERMINADA'}"),
                "tabla_documental": documented_table or document_name,
            })
    return result


def extract_pdf_table(path: Path, source: str, instrument: str,
                      member: str | None = None) -> list[dict]:
    """Reconstruye la columna Pregunta de diccionarios PDF multipágina.

    Las variables buscadas provienen del índice histórico del mismo instrumento;
    por ello un token ajeno al objeto nunca puede convertirse en correspondencia.
    """
    import pdfplumber

    variables, variable_tables = metadata_tables(instrument)
    if member:
        with zipfile.ZipFile(path) as archive:
            handle = io.BytesIO(archive.read(member))
        context = pdfplumber.open(handle)
        document_name = member
    else:
        context = pdfplumber.open(str(path))
        document_name = source

    known_tables = sorted({table for tables in variable_tables.values() for table in tables})
    current_table = ""
    current_question = ""
    dictionary_started = False
    result = []
    with context as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            page_text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            if "diccionario de datos de la base de datos por tabla" in fold(page_text):
                dictionary_started = True
            if not dictionary_started:
                continue
            for match in re.finditer(r"(?:^|\n)\s*(?:Tabla\s+|3\.\d+\s+)([A-Za-z][A-Za-z0-9_]+)", page_text):
                proposed = match.group(1)
                compatible = [table for table in known_tables if table_matches(table, proposed)]
                if compatible:
                    if proposed != current_table:
                        current_question = ""
                    current_table = proposed

            for table in page.extract_tables():
                rows = [[cell or "" for cell in row] for row in table if row]
                variable_column = None
                best_count = 0
                for row in rows:
                    for column, cell in enumerate(row):
                        count = sum(resolve_variable(token, variables) is not None
                                    for token in re.findall(r"[A-Za-z][A-Za-z0-9_]*", cell))
                        if count > best_count:
                            variable_column, best_count = column, count
                if variable_column is None:
                    continue
                question_column = max(0, variable_column - 1)
                for row in rows:
                    variable_ids = [resolved for token in re.findall(
                                        r"[A-Za-z][A-Za-z0-9_]*", row[variable_column])
                                    if (resolved := resolve_variable(token, variables)) is not None]
                    type_cell = (fold(row[variable_column + 1])
                                 if variable_column + 1 < len(row) else "")
                    has_type = valid_data_type(type_cell)
                    if variable_ids and not has_type:
                        variable_ids = []
                    raw_question = row[question_column] if question_column < len(row) else ""
                    question_cell = sanitize(raw_question)
                    if not variable_ids:
                        if (question_cell and fold(question_cell).strip(".: ")
                                not in {"pregunta", "(1)", "cons"}):
                            current_question = question_cell
                        continue

                    items = [sanitize(item) for item in re.split(r"\n(?=\s*\d+[.)]\s*)", raw_question)
                             if sanitize(item)]
                    if len(items) != len(variable_ids):
                        items = [question_cell] * len(variable_ids)
                    concept_column = variable_column + 3
                    concept = (sanitize(row[concept_column])
                               if concept_column < len(row) else "")
                    for variable_id, item_text in zip(variable_ids, items):
                        is_reactive = bool(re.match(r"^(?:[ab]?p)\d", variable_id, re.I))
                        if not is_reactive and concept:
                            question = concept
                        else:
                            question = clean_pdf_text(" ".join(dict.fromkeys(
                                part for part in (current_question, item_text) if part)))
                        text_type = "PREGUNTA_DICCIONARIO"
                        if not question and concept:
                            question = concept
                            text_type = "DESCRIPCION_ADMINISTRATIVA"
                        elif not ("?" in question or re.match(r"^\d+(?:\.\d+)+", question)):
                            text_type = "DESCRIPCION_ADMINISTRATIVA"
                        if not question:
                            continue
                        possible_tables = variable_tables[variable_id]
                        documented_table = current_table
                        if current_table and not any(table_matches(table, current_table)
                                                     for table in possible_tables):
                            documented_table = ""
                        if not documented_table:
                            documented_table = inferred_table(possible_tables)
                        result.append({
                            "variable_id": variable_id,
                            "texto_reactivo": question,
                            "texto_tipo": text_type,
                            "contexto_busqueda": "",
                            "fuente_texto": source,
                            "referencia_fuente": (
                                f"miembro={document_name};pagina={page_number};tabla={documented_table or 'NO_DETERMINADA'}"),
                            "tabla_documental": documented_table or document_name,
                        })
            covered = {(item["tabla_documental"], item["variable_id"]) for item in result}
            for item in extract_pdf_words(
                    page, variables, variable_tables, source, document_name,
                    page_number, current_table):
                if (item["tabla_documental"], item["variable_id"]) not in covered:
                    result.append(item)
                    covered.add((item["tabla_documental"], item["variable_id"]))
    return result


def extract_source(path: Path, row: dict) -> list[dict]:
    fmt = row["formato"]
    source = row["fuente_texto"]
    if fmt == "xlsx":
        return extract_xlsx(path, source)
    if fmt == "xls":
        return extract_xls(path, source)
    if fmt == "zip_xlsx":
        result = []
        with zipfile.ZipFile(path) as archive:
            for member in archive.namelist():
                if member.lower().endswith(".xlsx"):
                    result.extend(extract_xlsx(io.BytesIO(archive.read(member)), source, member))
        return result
    if fmt == "fd_pdf_index":
        return extract_fd_pdf_index(path, source, row["instrumento"])
    if fmt == "pdf_table":
        return extract_pdf_table(path, source, row["instrumento"])
    if fmt == "zip_pdf_table":
        result = []
        with zipfile.ZipFile(path) as archive:
            for member in archive.namelist():
                if member.lower().endswith(".pdf"):
                    result.extend(extract_pdf_table(path, source, row["instrumento"], member))
        return result
    raise ValueError(f"formato no soportado: {fmt}")


def cached_extract(path: Path, row: dict, cache_dir: Path) -> tuple[list[dict], bool, str]:
    digest = sha256_file(path)
    key_material = (f"{EXTRACTOR_VERSION}\0{row['instrumento']}\0"
                    f"{row['formato']}\0{digest}").encode()
    key = hashlib.sha256(key_material).hexdigest()
    cache_path = cache_dir / f"{key}.json"
    if cache_path.exists():
        return json.loads(cache_path.read_text(encoding="utf-8")), True, digest
    extracted = extract_source(path, row)
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(json.dumps(extracted, ensure_ascii=False, sort_keys=True), encoding="utf-8")
    return extracted, False, digest


def compatible_table(member: str, documented: str) -> bool:
    return table_matches(member, documented)


def choose_candidate(candidates: list[dict], member: str) -> dict | None:
    compatible = [item for item in candidates if compatible_table(member, item["tabla_documental"])]
    distinct = {(item["texto_reactivo"], item["fuente_texto"], item["referencia_fuente"]): item
                for item in compatible}
    if len(distinct) == 1:
        return next(iter(distinct.values()))
    same_text = {item["texto_reactivo"] for item in compatible}
    return compatible[0] if compatible and len(same_text) == 1 else None


def verified_candidates(path: Path, objects: list[str]) -> list[dict]:
    result = []
    for row in read_tsv(path):
        if not selected(row["instrumento"], row["fuente_texto"], objects):
            continue
        source_path = RAW_ROOT / row["fuente_texto"]
        if not source_path.exists():
            raise FileNotFoundError(source_path)
        reference = (f"pagina={row['pagina']};seccion={row['seccion']}"
                     if row["pagina"] else row["seccion"])
        result.append({
            **row,
            "tabla_documental": row["seccion"],
            "referencia_fuente": reference,
            "fuente_sha256_12": sha256_file(source_path)[:12],
            "manual": True,
        })
    return result


def write_output(path: Path, rows: list[dict], summary: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(
                "w", encoding="utf-8", newline="", dir=path.parent,
                prefix=f".{path.name}.", suffix=".tmp", delete=False) as handle:
            tmp_path = Path(handle.name)
            try:
                display_path = path.resolve().relative_to(REPO_ROOT).as_posix()
            except ValueError:
                display_path = path.name
            handle.write(f"# {display_path} -- sucesor overlay; no reordena fuentes historicas\n")
            handle.write("# contexto_busqueda es vocabulario editorial de recuperacion, no texto literal de la fuente\n")
            handle.write("# " + json.dumps(summary, ensure_ascii=False, sort_keys=True) + "\n")
            handle.write("\t".join(OUT_FIELDS) + "\n")
            for row in rows:
                handle.write("\t".join(sanitize(row.get(field, "")) for field in OUT_FIELDS) + "\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(tmp_path, path)
    finally:
        if tmp_path is not None and tmp_path.exists():
            tmp_path.unlink()


def write_residual(path: Path, rows: list[dict], summary: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    grouped: dict[tuple[str, ...], dict] = {}
    for row in rows:
        key = tuple(row[field] for field in
                    ("instrumento", "ola", "payload_id", "archivo_miembro", "motivo"))
        item = grouped.setdefault(key, {"count": 0, "variables": set(),
                                        "action": row["siguiente_accion"]})
        item["count"] += 1
        item["variables"].add(row["variable_id"])
    output = []
    for key, item in sorted(grouped.items()):
        variables = sorted(item["variables"], key=str.lower)
        output.append(dict(zip(
            ("instrumento", "ola", "payload_id", "archivo_miembro", "motivo"), key)) | {
                "filas_fisicas": item["count"],
                "variables_logicas": len(variables),
                "variables_ejemplo": ",".join(variables[:8]),
                "siguiente_accion": item["action"],
            })
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(
                "w", encoding="utf-8", newline="", dir=path.parent,
                prefix=f".{path.name}.", suffix=".tmp", delete=False) as handle:
            tmp_path = Path(handle.name)
            handle.write("# residual por objeto exacto; no equivale a ausencia cientifica\n")
            handle.write("# " + json.dumps(summary, ensure_ascii=False, sort_keys=True) + "\n")
            handle.write("\t".join(RESIDUAL_FIELDS) + "\n")
            for row in output:
                handle.write("\t".join(sanitize(row.get(field, ""))
                                       for field in RESIDUAL_FIELDS) + "\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(tmp_path, path)
    finally:
        if tmp_path is not None and tmp_path.exists():
            tmp_path.unlink()


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--objeto", action="append", help="Instrumento/familia o payload; repetible")
    parser.add_argument("--salida", type=Path, default=OUT_PATH)
    parser.add_argument("--fuentes", type=Path, default=SOURCES_PATH)
    parser.add_argument("--verificados", type=Path, default=VERIFIED_PATH)
    parser.add_argument("--cache-dir", type=Path, default=CACHE_DIR)
    parser.add_argument("--reporte-residual", type=Path, default=RESIDUAL_PATH)
    args = parser.parse_args(argv)
    objects = args.objeto or list(PRIORITY)

    metadata = []
    for source_key, path in METADATA_SOURCES.items():
        for position, row in enumerate(read_tsv(path), 1):
            instrument = instrument_for(row)
            if selected(instrument, row["payload_id"], objects):
                metadata.append((source_key, position, instrument, row))

    # La sucesion es aditiva: una identidad ya acreditada conserva literalmente
    # texto, tipo, contexto y cita. El extractor nuevo solo llena puntos ciegos.
    previous_by_origin = {
        row["id_origen"]: row for row in read_tsv(PREVIOUS_OVERLAY_PATH)
    }

    candidates: dict[tuple[str, str], list[dict]] = defaultdict(list)
    cache_hits = cache_misses = 0
    source_hashes = {}
    source_specs = read_tsv(args.fuentes)
    configured_instruments = {row["instrumento"].lower() for row in source_specs}
    for source in source_specs:
        if not selected(source["instrumento"], source["fuente_texto"], objects):
            continue
        path = RAW_ROOT / source["fuente_texto"]
        if not path.exists():
            continue
        extracted, hit, digest = cached_extract(path, source, args.cache_dir)
        cache_hits += int(hit)
        cache_misses += int(not hit)
        source_hashes[source["fuente_texto"]] = digest[:12]
        for item in extracted:
            item["fuente_sha256_12"] = digest[:12]
            candidates[(source["instrumento"].lower(), item["variable_id"].lower())].append(item)

    for item in verified_candidates(args.verificados, objects):
        key = (item["instrumento"].lower(), item["variable_id"].lower())
        candidates[key].insert(0, item)

    output = []
    residual_rows = []
    unresolved = 0
    for source_key, position, instrument, row in metadata:
        origin = f"{source_key}:{position}"
        native_text = sanitize(row["texto_reactivo"])
        candidate = None
        if origin in previous_by_origin:
            candidate = previous_by_origin[origin]
        elif native_text:
            candidate = {
                "texto_reactivo": native_text,
                "texto_tipo": "ETIQUETA_VARIABLE",
                "contexto_busqueda": "",
                "fuente_texto": row["payload_id"],
                "fuente_sha256_12": row["sha256_12"],
                "referencia_fuente": f"miembro={row['archivo_miembro']};metadato_nativo",
            }
        else:
            pool = candidates.get((instrument.lower(), row["variable_id"].lower()), [])
            manual = [item for item in pool if item.get("manual")]
            candidate = choose_candidate(
                manual if manual else pool, row["archivo_miembro"])
        if candidate is None:
            unresolved += 1
            member_folded = fold(row["archivo_miembro"])
            variable = row["variable_id"]
            if pool:
                reason = "CORRESPONDENCIA_AMBIGUA"
                action = "resolver tabla/miembro o conservar ausencia"
            elif any(token in member_folded for token in
                     ("/catalogos/", "/metadatos/", "/diccionario_de_datos/")):
                reason = "FILA_AUXILIAR_SIN_REACTIVO"
                action = "acreditar etiqueta tecnica del descriptor; no inventar pregunta"
            elif not re.match(r"^(?:[ab]?p)\d", variable, re.I):
                reason = "ETIQUETA_TECNICA_NO_ACREDITADA"
                action = "localizar etiqueta de llave/ponderador en el diccionario exacto"
            elif instrument.lower() not in configured_instruments:
                reason = "DOCUMENTO_DE_VARIABLE_NO_REGISTRADO"
                action = "registrar cuestionario/diccionario oficial de esta ola y tabla"
            else:
                reason = "PREGUNTA_NO_LOCALIZADA"
                action = "verificar el PDF/XLS exacto; OCR local si el glifo no es extraible"
            residual_rows.append({
                "instrumento": instrument,
                "ola": year_for(instrument, row["ola"]),
                "payload_id": row["payload_id"],
                "archivo_miembro": row["archivo_miembro"],
                "variable_id": variable,
                "motivo": reason,
                "siguiente_accion": action,
            })
            continue
        output.append({
            "id_origen": origin,
            "payload_id": row["payload_id"],
            "sha256_12": row["sha256_12"],
            "instrumento": instrument,
            "ola": year_for(instrument, row["ola"]),
            "archivo_miembro": row["archivo_miembro"],
            "variable_id": row["variable_id"],
            "metodo": row["metodo"],
            "universo_declarado": row["universo_declarado"],
            **{key: candidate.get(key, "") for key in (
                "texto_reactivo", "texto_tipo", "contexto_busqueda", "fuente_texto",
                "fuente_sha256_12", "referencia_fuente")},
        })

    output.sort(key=lambda row: (row["id_origen"].split(":")[0], int(row["id_origen"].split(":")[1])))
    summary = {
        "extractor_version": EXTRACTOR_VERSION,
        "objetos": objects,
        "filas_metadato_perimetro": len(metadata),
        "filas_con_texto_publicadas": len(output),
        "filas_sin_texto_residual": unresolved,
        "cache_hits": cache_hits,
        "cache_misses": cache_misses,
        "fuentes": source_hashes,
    }
    reason_counts = Counter(row["motivo"] for row in residual_rows)
    summary["grupos_residuales"] = len({(
        row["instrumento"], row["ola"], row["payload_id"],
        row["archivo_miembro"], row["motivo"]) for row in residual_rows})
    summary["motivos_residuales"] = dict(sorted(reason_counts.items()))
    write_output(args.salida, output, summary)
    write_residual(args.reporte_residual, residual_rows, summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
