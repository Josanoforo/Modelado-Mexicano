#!/usr/bin/env python3
"""Control independiente de unidad, diseño, varianza e IC de ENADID.

No importa ``medidor.py`` ni reutiliza sus funciones. Lee TSDEM con el módulo
CSV estándar, reconstruye los dos productos y contrasta la tabla publicada
sin modificar ningún archivo del CALC sellado.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import subprocess
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from scipy.stats import t as student_t


EXPECTED = {
    "datos": "995d448ab298251d441e437fc4a368f84879a8b74fa7bc103dc189434cbfaa18",
    "fd": "d56582b7ace04e2c13a1b9ea22aeea458d7b9cc02a21241e5fda29fe12b4f34c",
    "cuestionario": "8b046a68dc19e292522557e05d81952e8fb04c2d3a1375c93fa7192e857e3270",
}
AGES = (
    ("15_17", 15, 17), ("18_29", 18, 29), ("30_44", 30, 44),
    ("45_59", 45, 59), ("60_mas", 60, 120), ("15_mas", 15, 120),
)
VALID_STATUS = set("1234567")
TOLERANCE = 5e-10


def sha256(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def documentary_checks(fd_path: str, questionnaire_path: str) -> dict:
    workbook = load_workbook(fd_path, read_only=True, data_only=True)
    sheet = workbook["TSDEM"]
    fields = {}
    current = None
    for row in sheet.iter_rows(values_only=True):
        mnemonic = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        if mnemonic:
            current = mnemonic.upper()
            fields.setdefault(current, {"description": row[4], "valid": row[7], "classes": []})
        if current and len(row) > 9 and row[8] is not None:
            fields[current]["classes"].append((str(row[8]).strip(), str(row[9]).strip()))
    required = {"LLAVE_PER", "EDAD", "P3_27", "P3_27_AG", "FAC_VIV", "EST_DIS", "UPM_DIS"}
    missing = sorted(required - set(fields))
    p327 = dict(fields["P3_27"]["classes"])
    expected_labels = {
        "1": "vive con su pareja en unión libre?",
        "2": "está separada(o) de una unión libre?",
        "3": "está separada(o) de un matrimonio?",
        "4": "está divorciada(o)?",
        "5": "esta viuda(o)?",
        "6": "está casada(o)?",
        "7": "está soltera(o)?",
    }
    questionnaire = subprocess.run(
        ["pdftotext", "-layout", questionnaire_path, "-"],
        check=True, capture_output=True, text=True,
    ).stdout
    return {
        "hoja": "TSDEM",
        "campos_requeridos_presentes": not missing,
        "campos_faltantes": missing,
        "p3_27_codigos_y_etiquetas_ok": p327 == expected_labels,
        "p3_27_ag_identificada_como_agrupada": "agrupada" in str(fields["P3_27_AG"]["description"]).lower(),
        "fac_viv_es_factor_en_tsdem": "Factor de expansión" in str(fields["FAC_VIV"]["description"]),
        "est_dis_es_estrato_diseno": "diseño muestral" in str(fields["EST_DIS"]["description"]),
        "upm_dis_es_upm_diseno": "Unidad Primaria de Muestreo" in str(fields["UPM_DIS"]["description"]),
        "cuestionario_universo_12_mas": "PARA PERSONAS DE 12 AÑOS CUMPLIDOS O MÁS" in questionnaire,
        "cuestionario_situacion_actual": "¿Actualmente (NOMBRE)..." in questionnaire,
        "cuestionario_union_libre": "vive con su pareja en unión libre?" in questionnaire,
        "cuestionario_casada": "está casada(o)?" in questionnaire,
    }


def load_rows(zip_path: str) -> tuple[list[tuple], dict]:
    rows = []
    keys = set()
    duplicate_keys = 0
    invalid_status = invalid_weight = missing_design = age_unknown = 0
    with zipfile.ZipFile(zip_path) as archive, archive.open("TSDEM.csv") as raw:
        reader = csv.DictReader(io.TextIOWrapper(raw, encoding="utf-8-sig", newline=""))
        for source in reader:
            row = {key.upper(): value.strip() for key, value in source.items()}
            key = row["LLAVE_PER"]
            if key in keys:
                duplicate_keys += 1
            keys.add(key)
            try:
                age = int(row["EDAD"])
            except ValueError:
                age = None
            try:
                weight = float(row["FAC_VIV"])
            except ValueError:
                weight = math.nan
            weight_ok = math.isfinite(weight) and weight > 0
            design_ok = bool(row["EST_DIS"] and row["UPM_DIS"])
            if row["P3_27"] not in VALID_STATUS:
                invalid_status += 1
            if not weight_ok:
                invalid_weight += 1
            if not design_ok:
                missing_design += 1
            if age is None or not 0 <= age <= 120:
                age_unknown += 1
            rows.append((key, age, row["P3_27"], weight, weight_ok,
                         row["EST_DIS"], row["UPM_DIS"], design_ok))
    known_15 = [row for row in rows if row[1] is not None and 15 <= row[1] <= 120]
    status_counts_15 = {code: sum(row[2] == code for row in known_15) for code in sorted(VALID_STATUS)}
    return rows, {
        "filas": len(rows), "llaves_unicas": len(keys),
        "llaves_duplicadas": duplicate_keys, "situacion_fuera_catalogo_total_n": invalid_status,
        "peso_invalido_n": invalid_weight, "diseno_faltante_total_n": missing_design,
        "edad_no_especificada_n": age_unknown,
        "edad_conocida_menor_15_n": sum(row[1] is not None and 0 <= row[1] < 15 for row in rows),
        "edad_conocida_15_mas_n": len(known_15),
        "situacion_invalida_15_mas_n": sum(row[2] not in VALID_STATUS for row in known_15),
        "situacion_codigos_15_mas": status_counts_15,
        "fuera_denominador_condicional_15_mas_n": sum(row[2] in VALID_STATUS - {"1", "6"} for row in known_15),
        "diseno_faltante_15_mas_n": sum(not row[7] for row in known_15),
    }


def estimate(rows: list[tuple], low: int, high: int, conditional: bool) -> dict:
    def masks(row):
        _, age, status, _, weight_ok, _, _, _ = row
        age_ok = age is not None and low <= age <= high
        valid = status in ({"1", "6"} if conditional else VALID_STATUS)
        return age_ok and valid and weight_ok, age_ok and status == "1" and weight_ok

    denominator_rows = [row for row in rows if masks(row)[0]]
    numerator_rows = [row for row in rows if masks(row)[1]]
    denominator = math.fsum(row[3] for row in denominator_rows)
    numerator = math.fsum(row[3] for row in numerator_rows)
    point = numerator / denominator

    psu_totals = defaultdict(float)
    design_missing_domain = 0
    for row in rows:
        _, _, _, weight, weight_ok, stratum, psu, design_ok = row
        x, y = masks(row)
        if x and not design_ok:
            design_missing_domain += 1
        if not weight_ok or not design_ok:
            continue
        influence = weight * (float(y) - point * float(x)) / denominator
        psu_totals[(stratum, psu)] += influence

    strata = defaultdict(list)
    for (stratum, _), value in psu_totals.items():
        strata[stratum].append(value)
    terms = []
    singletons = 0
    df = 0
    for values in strata.values():
        count = len(values)
        if count == 1:
            singletons += 1
            continue
        mean = math.fsum(values) / count
        terms.append(count / (count - 1) * math.fsum((value - mean) ** 2 for value in values))
        df += count - 1
    variance = math.fsum(terms) * (len(terms) + singletons) / len(terms)
    se = math.sqrt(variance)
    critical = float(student_t.ppf(0.975, df))
    logit = math.log(point / (1 - point))
    se_logit = se / (point * (1 - point))
    inverse = lambda value: 1 / (1 + math.exp(-value))
    return {
        "n": len(denominator_rows), "numerator_n": len(numerator_rows),
        "denominator": denominator, "numerator": numerator, "point": point,
        "se": se, "lo": inverse(logit - critical * se_logit),
        "hi": inverse(logit + critical * se_logit), "df": df,
        "strata": len(strata), "singletons": singletons,
        "psu_nested": len(psu_totals), "design_missing_domain": design_missing_domain,
    }


def compare(zip_path: str, fd_path: str, questionnaire_path: str, published_path: str) -> dict:
    identities = {
        "datos": sha256(zip_path), "fd": sha256(fd_path),
        "cuestionario": sha256(questionnaire_path),
    }
    rows, audit = load_rows(zip_path)
    with open(published_path, encoding="utf-8", newline="") as handle:
        published = list(csv.DictReader(handle))
    checks = []
    for age_id, low, high in AGES:
        for conditional in (False, True):
            estimand = "union_libre_entre_union_o_casada" if conditional else "distribucion_actual_15_mas"
            expected = next(row for row in published if row["estimando"] == estimand and row["categoria"] == "union_libre" and row["edad"] == age_id)
            actual = estimate(rows, low, high, conditional)
            deltas = {
                "punto": abs(actual["point"] - float(expected["punto"])),
                "ee": abs(actual["se"] - float(expected["ee"])),
                "ic95_lo": abs(actual["lo"] - float(expected["ic95_lo"])),
                "ic95_hi": abs(actual["hi"] - float(expected["ic95_hi"])),
            }
            checks.append({
                "producto": "condicional_union_actual" if conditional else "bruto_total_elegible",
                "edad": age_id, "n_publicado": int(expected["n_valido"]),
                "n_control": actual["n"], "df_publicado": int(expected["gl"]),
                "df_control": actual["df"], "estratos_control": actual["strata"],
                "upm_control": actual["psu_nested"], "singleton_control": actual["singletons"],
                "diseno_faltante_dominio_n": actual["design_missing_domain"],
                "deltas": deltas,
                "ok": (actual["n"] == int(expected["n_valido"])
                       and actual["df"] == int(expected["gl"])
                       and all(delta <= TOLERANCE for delta in deltas.values())),
            })
    documentary = documentary_checks(fd_path, questionnaire_path)
    return {
        "metodo": "csv/openpyxl/scipy independientes; no importa medidor.py",
        "identidades": identities,
        "identidades_ok": identities == EXPECTED,
        "unidad_estimacion": "persona residente; una fila por LLAVE_PER en TSDEM",
        "productos": {
            "bruto_total_elegible": "P3_27=1 / P3_27 en 1..7, edad conocida del dominio",
            "condicional_union_actual": "P3_27=1 / P3_27 en {1,6}, edad conocida del dominio",
        },
        "productos_distintos": True,
        "auditoria_unidad_soporte": audit,
        "evidencia_documental": documentary,
        "politica_singleton": "aporte medio de estratos no singleton; singleton sin gl propio",
        "fpc": "no aplicada; no acreditada",
        "tolerancia_absoluta": TOLERANCE,
        "checks": checks,
        "todos_ok": (identities == EXPECTED and audit["llaves_duplicadas"] == 0
                     and all(value is True for key, value in documentary.items()
                             if key.endswith("_ok") or key.startswith("cuestionario_")
                             or key.endswith("_presente") or key.endswith("_diseno")
                             or key.endswith("_tsdem") or key.endswith("_agrupada"))
                     and all(check["ok"] for check in checks)),
    }


def main(arguments: list[str]) -> int:
    if len(arguments) != 5:
        raise SystemExit("uso: control_precision_independiente.py DATOS FD CUESTIONARIO RESULTADOS SALIDA")
    result = compare(*arguments[:4])
    Path(arguments[4]).write_text(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return 0 if result["todos_ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
