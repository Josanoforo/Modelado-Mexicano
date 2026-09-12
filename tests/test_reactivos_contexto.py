import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from tools import actualiza_reactivos_contexto as arc
from tools import busca_reactivos as busca


class ReactivosContextoTest(unittest.TestCase):
    def test_extract_sheet_uses_question_from_preceding_row(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "TMod_Vic"
        sheet.append(["Pregunta", "Nemónico"])
        sheet.append(["1.23 ¿Por qué no denunció?", None])
        sheet.append([None, "BP1_23"])
        rows = arc.extract_sheet(sheet, "fd.pdf", "fd.pdf")
        self.assertEqual("BP1_23", rows[0]["variable_id"])
        self.assertEqual("1.23 ¿Por qué no denunció?", rows[0]["texto_reactivo"])
        self.assertIn("fila=3", rows[0]["referencia_fuente"])

    def test_extract_sheet_keeps_question_for_all_subitems(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Pregunta", "Nemónico"])
        sheet.append(["5.1 ¿Cuáles productos tiene?", None])
        sheet.append([None, "P5_1_1"])
        sheet.append([None, "P5_1_2"])
        rows = arc.extract_sheet(sheet, "fd.xlsx", "fd.xlsx")
        self.assertEqual(["P5_1_1", "P5_1_2"], [row["variable_id"] for row in rows])
        self.assertEqual(
            ["5.1 ¿Cuáles productos tiene?"] * 2,
            [row["texto_reactivo"] for row in rows],
        )

    def test_cache_key_skips_unchanged_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "fd.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Pregunta", "Nemónico"])
            sheet.append(["¿Ahorra?", "P1"])
            workbook.save(source)
            spec = {"formato": "xlsx", "fuente_texto": "fd.xlsx", "instrumento": "enif2024"}
            first, first_hit, digest1 = arc.cached_extract(source, spec, root / "cache")
            second, second_hit, digest2 = arc.cached_extract(source, spec, root / "cache")
            self.assertFalse(first_hit)
            self.assertTrue(second_hit)
            self.assertEqual(first, second)
            self.assertEqual(digest1, digest2)

    def test_pdf_cleanup_drops_layout_footer(self):
        dirty = "4.10 ¿Dejó de tomar taxi? (Continúa) INEGI. Encuesta Nacional de Victimización"
        self.assertEqual("4.10 ¿Dejó de tomar taxi?", arc.clean_pdf_text(dirty))
        self.assertTrue(arc.valid_data_type("Alfanumérica"))

    def test_repeated_verified_variable_is_selected_by_exact_table(self):
        candidates = [
            {"tabla_documental": "THogar", "texto_reactivo": "hogar",
             "fuente_texto": "fd.pdf", "referencia_fuente": "pagina=12"},
            {"tabla_documental": "TMod_Vic", "texto_reactivo": "delito",
             "fuente_texto": "fd.pdf", "referencia_fuente": "pagina=52"},
        ]
        chosen = arc.choose_candidate(candidates, "ruta/TMod_Vic.dbf")
        self.assertEqual("delito", chosen["texto_reactivo"])

    def test_table_correspondence_distinguishes_repeated_variables(self):
        self.assertTrue(arc.table_matches("carpeta/TPer_Vic1.dbf", "TPer_Vic1"))
        self.assertFalse(arc.table_matches("carpeta/TPer_Vic1.dbf", "TPer_Vic2"))
        candidates = [{
            "texto_reactivo": "texto de otra tabla",
            "fuente_texto": "fd.pdf",
            "referencia_fuente": "pagina=1;tabla=THogar",
            "tabla_documental": "THogar",
        }]
        self.assertIsNone(arc.choose_candidate(candidates, "TSDem.dbf"))

    def test_ennvih_instrument_and_wave_are_derived_from_payload(self):
        row = {"instrumento": "(sin-instrumento-derivable)", "payload_id": "ennvih/ehh05dta_all.zip"}
        self.assertEqual("ennvih2005", arc.instrument_for(row))
        self.assertEqual("2005", arc.year_for("ennvih2005", "NO_DETERMINADO"))

    def test_output_schema_is_tsv_safe(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.tsv"
            row = {field: "" for field in arc.OUT_FIELDS}
            row.update({"id_origen": "v1_2:1", "texto_reactivo": "línea\ncon\ttabs"})
            arc.write_output(path, [row], {"ok": True})
            with path.open(encoding="utf-8") as handle:
                parsed = list(csv.DictReader((line for line in handle if not line.startswith("#")), delimiter="\t"))
            self.assertEqual("línea con tabs", parsed[0]["texto_reactivo"])

    def test_search_filter_uses_accredited_context(self):
        args = type("Args", (), {
            "palabra": ["corrupcion"], "regex": None, "encuesta": "encuci",
            "ola": None, "tipo": None,
        })()
        predicate = busca.construye_filtro(args)
        row = {
            "instrumento": "encuci2020", "ola": "2020", "payload_id": "bd.zip",
            "metodo": "INSPECT_ZIP", "texto_reactivo": "¿Le pidieron una dádiva?",
            "variable_id": "AP5_17", "contexto_busqueda": "Sección corrupción",
        }
        self.assertTrue(predicate(row))

    def test_published_overlay_preserves_origins_and_verified_cases(self):
        rows = arc.read_tsv(arc.OUT_PATH)
        previous = arc.read_tsv(arc.PREVIOUS_OVERLAY_PATH)
        self.assertGreater(len(rows), len(previous))
        self.assertEqual(len(rows), len({row["id_origen"] for row in rows}))
        self.assertTrue(all(row["texto_reactivo"] and row["fuente_texto"] and
                            row["fuente_sha256_12"] and row["referencia_fuente"]
                            for row in rows))

        by_case = {(row["instrumento"].lower(), row["variable_id"].lower()): row
                   for row in rows}
        expected = {
            ("envipe2025", "bp1_23"),
            ("ennvih2002", "crh01_1e"),
            ("enif2024", "p5_1_5"),
            ("encuci2020", "ap5_17"),
            ("ensafi2023", "p6_7"),
            ("ensafi2023", "p6_10_2"),
        }
        self.assertTrue(expected <= by_case.keys())

        by_origin = {row["id_origen"]: row for row in rows}
        inherited_fields = (
            "texto_reactivo", "texto_tipo", "contexto_busqueda", "fuente_texto",
            "fuente_sha256_12", "referencia_fuente",
        )
        for old in previous:
            self.assertIn(old["id_origen"], by_origin)
            self.assertEqual(
                tuple(old[field] for field in inherited_fields),
                tuple(by_origin[old["id_origen"]][field] for field in inherited_fields),
            )

        historical = {}
        for source, path in arc.METADATA_SOURCES.items():
            for position, row in enumerate(arc.read_tsv(path), 1):
                historical[f"{source}:{position}"] = row
        for row in rows:
            origin = historical[row["id_origen"]]
            self.assertEqual(
                (row["payload_id"], row["archivo_miembro"], row["variable_id"], row["sha256_12"]),
                (origin["payload_id"], origin["archivo_miembro"], origin["variable_id"],
                 origin["sha256_12"]),
            )


if __name__ == "__main__":
    unittest.main()
