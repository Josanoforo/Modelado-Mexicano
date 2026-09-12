"""Control material independiente del lector XML y del estimador principal."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook


PRODUCTS = {
    "tdc": "tarjeta_credito", "hip": "credito_hipotecario",
    "per": "credito_personal", "nom": "credito_nomina",
    "aut": "credito_automotriz",
}


def code(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def ratio(rows, col, allowed, positive):
    valid = [row for row in rows if code(row[col]) in allowed]
    selected = [row for row in valid if code(row[col]) in positive]
    denominator = math.fsum(float(row["ponderador"]) for row in valid)
    numerator = math.fsum(float(row["ponderador"]) for row in selected)
    return {
        "n_elegible": len(rows), "n_valido": len(valid),
        "n_positivo": len(selected), "masa_valida": denominator,
        "masa_positiva": numerator, "estimacion": numerator / denominator,
    }


def close(a, b, tol=1e-10):
    return abs(float(a) - float(b)) <= tol


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, default=Path("data/raw/GEN2_N34_PRODUCTO_DANO/banxico_satisfaccion_usuarios_2019_2024.xlsx"))
    parser.add_argument("--outdir", type=Path, default=Path("data/banxico-producto-dano-medicion"))
    parser.add_argument("--coverage", type=Path, default=Path("data/n34-producto-dano/banxico-cobertura-producto-ola.csv"))
    parser.add_argument("--results", type=Path, default=Path("data/corrida0/CALC-BANXICO-PRODUCTO-DANO-0001/resultados.json"))
    parser.add_argument("--write", type=Path, default=Path("data/corrida0/CALC-BANXICO-PRODUCTO-DANO-0001/control-independiente.json"))
    args = parser.parse_args()

    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    header = [code(v) for v in next(iterator)]
    rows = [dict(zip(header, values)) for values in iterator]
    workbook.close()

    with (args.outdir / "estimandos.csv").open(encoding="utf-8", newline="") as handle:
        estimates = list(csv.DictReader(handle))

    def artifact_row(kind, product, wave):
        matches = [r for r in estimates if r["estimando_tipo"] == kind and r["producto"] == product and r["ola"] == str(wave)]
        if len(matches) != 1:
            raise AssertionError(f"fila no univoca: {kind}/{product}/{wave}: {len(matches)}")
        return matches[0]

    tdc_2024 = [r for r in rows if int(r["fecha"]) == 2024 and code(r["tdc_filtro"]) == "1"]
    tdc = ratio(tdc_2024, "tdc_comp_pago", {"1", "2", "3", "4"}, {"3", "4"})
    tdc_out = artifact_row("dano_combinado", "tarjeta_credito", 2024)
    tdc_ok = all((int(tdc_out[k]) == tdc[k]) for k in ("n_elegible", "n_valido", "n_positivo")) and close(tdc_out["masa_valida"], tdc["masa_valida"], 5e-6) and close(tdc_out["masa_positiva"], tdc["masa_positiva"], 5e-6) and close(tdc_out["estimacion"], tdc["estimacion"])

    hip_2024 = [r for r in rows if int(r["fecha"]) == 2024 and code(r["hip_filtro"]) == "1" and code(r["hip_problemas"]) == "1"]
    hip_claim = ratio(hip_2024, "hip_reclamacion", {"1", "2"}, {"1"})
    hip_out = artifact_row("reclamacion", "credito_hipotecario", 2024)
    hip_ok = all((int(hip_out[k]) == hip_claim[k]) for k in ("n_elegible", "n_valido", "n_positivo")) and close(hip_out["estimacion"], hip_claim["estimacion"])

    hip_2019 = [r for r in rows if int(r["fecha"]) == 2019 and code(r["hip_filtro"]) == "1"]
    hip_2019_codes = {code(r["hip_comp_pago"]) for r in hip_2019}
    no_raised_out = artifact_row("dano_combinado", "credito_hipotecario", 2019)
    no_raised_ok = hip_2019_codes <= {"", "-1"} and int(no_raised_out["n_valido"]) == 0 and no_raised_out["estimacion"] == "" and f"no_levantada={len(hip_2019)}" in no_raised_out["faltantes_por_causa"]

    with args.coverage.open(encoding="utf-8", newline="") as handle:
        coverage = {(r["producto"], int(r["periodo"])): r for r in csv.DictReader(handle)}
    coverage_checks = []
    for pfx, product in PRODUCTS.items():
        holders = [r for r in rows if int(r["fecha"]) == 2024 and code(r[f"{pfx}_filtro"]) == "1"]
        mass = math.fsum(float(r["ponderador"]) for r in holders)
        reference = coverage[(product, 2024)]
        ok = len(holders) == int(reference["n_tenedores"]) and close(mass, reference["masa_ponderada_tenedores"], 5e-7)
        coverage_checks.append({"producto": product, "n": len(holders), "masa": mass, "delta_masa": mass - float(reference["masa_ponderada_tenedores"]), "coincide": ok})

    sealed = json.loads(args.results.read_text(encoding="utf-8"))["resultados"]
    hash_map = {
        "estimandos.csv": "RESULT-BANXICO-G-ESTIMANDOS-SHA256",
        "conjunta-costo-pago.csv": "RESULT-BANXICO-G-CONJUNTA-SHA256",
        "costo-dano-2024.csv": "RESULT-BANXICO-G-COSTO-DANO-2024-SHA256",
        "resumen-producto-ola.csv": "RESULT-BANXICO-G-RESUMEN-SHA256",
        "atraso-por-producto-ola.svg": "RESULT-BANXICO-G-GRAFICO-SHA256",
    }
    hashes = []
    for name, result_id in hash_map.items():
        actual = hashlib.sha256((args.outdir / name).read_bytes()).hexdigest()
        hashes.append({"archivo": name, "sha256": actual, "result_id": result_id, "coincide": actual == sealed[result_id]})

    checks = {
        "cociente_tdc_dano_2024": {**tdc, "coincide": tdc_ok},
        "filtro_reclamacion_hipotecario_2024": {**hip_claim, "coincide": hip_ok},
        "no_levantado_hipotecario_2019": {"n_tenedores": len(hip_2019), "codigos_observados": sorted(hip_2019_codes), "coincide": no_raised_ok},
        "cobertura_2024_vs_734": coverage_checks,
        "hashes_artefactos_vs_result": hashes,
    }
    ok = tdc_ok and hip_ok and no_raised_ok and all(x["coincide"] for x in coverage_checks) and all(x["coincide"] for x in hashes)
    document = {"calc_id": "CALC-BANXICO-PRODUCTO-DANO-0001", "lector": "openpyxl-3.1.5; independiente del lector XML stdlib", "estado": "COINCIDE" if ok else "DISCREPA", "comprobaciones": checks}
    args.write.write_text(json.dumps(document, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(f"CONTROL-INDEPENDIENTE: {document['estado']} · 3 riesgos materiales · cobertura={sum(x['coincide'] for x in coverage_checks)}/5 · hashes={sum(x['coincide'] for x in hashes)}/5")
    raise SystemExit(0 if ok else 1)


if __name__ == "__main__":
    main()
