"""Transporta sólo diez documentos públicos allowlisted; nunca abre ZIP de datos."""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import subprocess
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--custodia', required=True)
    args = parser.parse_args()
    out = Path(args.custodia).resolve()
    if out == ROOT or ROOT in out.parents:
        raise ValueError('custodia debe quedar fuera del clon')
    config = json.loads((HERE / 'impedimentos-lote2-p4-documentos-publicos-contrato.json').read_text())
    rows, evidence = [], []
    for doc in config['documentos']:
        source = ROOT / doc['ruta_fuente']
        blob = source.read_bytes()
        if hashlib.sha256(blob).hexdigest() != doc['sha256']:
            raise ValueError('hash documental incorrecto: ' + doc['id'])
        if doc['tipo'].startswith('XLSX'):
            if not blob.startswith(b'PK'):
                raise ValueError('XLSX inválido')
            wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
            expected = str(doc['ola'])
            if not any(expected in name for name in wb.sheetnames):
                raise ValueError('ola FD no acreditada')
            variables = {}
            for sheet in wb:
                for n, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    cells = [str(v).strip() for v in row if v is not None]
                    for value in cells:
                        if value in {'UPM','VIV_SEL','HOGAR','NUM_REN','EDAD','SEXO','NIVEL','EST_DIS','UPM_DIS','FAC_PER','FAC_MOCIBA','ENT','CVE_ENT','TLOC','P7_1','P7_2','P7_10_2','P7_12_3','P7_35_4','P8_1','P8_2','P1_1','P1_10','P7_1_1','P7_1_6A','P4_01','P4_10','P10_1','P10_5'}:
                            variables[value] = {'hoja':sheet.title,'fila':n,'texto':' | '.join(cells)}
            required = {'EDAD','SEXO','NIVEL','EST_DIS','UPM_DIS'}
            required |= {'FAC_PER','P7_1','P7_10_2','P8_1'} if doc['instrumento']=='endutih' else {'FAC_MOCIBA'}
            missing = sorted(required - variables.keys())
            if missing:
                raise ValueError('variables FD faltantes: ' + str(missing))
            evidence.append({'documento':doc['id'],'hojas':wb.sheetnames,'variables':variables,'dictamen':'FD exacto por ola; reactivos/diseño documentados, sin registros observados'})
        else:
            if not blob.startswith(b'%PDF'):
                raise ValueError('PDF inválido')
            proc = subprocess.run(['pdftotext','-layout',str(source),'-'],check=True,capture_output=True,text=True)
            text = proc.stdout
            if str(doc['ola']) not in text[:4500] or ('CIBERACOSO' not in text[:4500].upper() if doc['instrumento']=='mociba' else 'HOGARES' not in text[:4500].upper()):
                raise ValueError('instrumento/ola PDF no acreditados')
            terms = ['15 años','tres meses','7.10','7.1','8.1'] if doc['instrumento']=='endutih' else (['12 años','doce meses','Policía','Bloquear'] if doc['ola']==2016 else ['12 a 59','junio de 2016','proveedor','Bloquear'])
            found = {term:[i+1 for i,line in enumerate(text.splitlines()) if term.casefold() in line.casefold()] for term in terms}
            evidence.append({'documento':doc['id'],'titulo':text[:500],'terminos_lineas':found,'dictamen':'formulario vacío exacto; ventanas/filtros/acciones localizados; no resultados'})
        target = out / (doc['id'] + ('.xlsx' if doc['tipo'].startswith('XLSX') else '.pdf'))
        rows.append(dict(doc, salida=str(target), bytes=len(blob)))
    # Verificar todo antes de escribir; salida exclusiva documental.
    out.mkdir(parents=True,exist_ok=True)
    for row in rows:
        target = Path(row['salida'])
        if target.is_symlink():
            raise ValueError('salida enlazada rechazada')
        target.write_bytes((ROOT / row['ruta_fuente']).read_bytes())
    (out/'recibo-documentos.json').write_text(json.dumps({'documentos':rows,'revision_semantica':evidence},ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'estado':'PASS','documentos':len(rows),'custodia':str(out),'microdatos_abiertos':0}))

if __name__ == '__main__':
    main()
